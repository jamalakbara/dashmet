"""PowerPoint (.pptx) export of a single-account overview.

Pure builder: `generate_overview_pptx` takes the dicts already produced by the
shared insights read functions (`get_overview` / `get_timeseries` / `get_table`)
and lays them onto a branded deck. It performs NO metric computation and NO DB
access — it only reads and formats values the read path already resolved
(P-6/P-7: reports use the same numbers as the API by construction, and there is
exactly one code path that computes a metric).

The visual language mirrors a monthly-report agency template (dark gradient
cover, indigo section headers, "Current vs Previous" hero, gray metric cards
with coloured deltas, editable purple insight boxes, current-vs-previous trend
overlay, closing slide) but with neutral dashmet branding — no third-party
logos.

The only network I/O is `_fetch_thumbnail`, a best-effort image download for the
top-ads slide. A failed/missing thumbnail degrades to a neutral placeholder — it
never fakes a metric and never fails the whole export.
"""

from __future__ import annotations

import io
import logging
import os
from datetime import date, datetime, timezone
from typing import Optional

import httpx
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

logger = logging.getLogger(__name__)

# ─── Canvas (16:9) ───────────────────────────────────────────────────────────────
# python-pptx measures in EMU; 914400 EMU = 1 inch.
_SLIDE_W = Emu(12192000)   # 13.333"
_SLIDE_H = Emu(6858000)    # 7.5"
_MARGIN = Emu(560000)

# ─── Palette (BDD-style, generic dashmet branding) ───────────────────────────────
_COVER_BG = RGBColor(0x0E, 0x11, 0x16)   # near-black cover/divider/closing
_INK = RGBColor(0x22, 0x25, 0x2B)        # primary text on light
_MUTED = RGBColor(0x8A, 0x90, 0x9A)      # secondary text
_INDIGO = RGBColor(0x3E, 0x3A, 0x6E)     # headers, table header row
_PURPLE_BOX = RGBColor(0x4A, 0x45, 0x70)  # insight box
_CARD = RGBColor(0xF2, 0xF3, 0xF5)       # metric card fill
_CARD_ALT = RGBColor(0xFB, 0xFB, 0xFC)   # zebra row
_HAIRLINE = RGBColor(0xE6, 0xE8, 0xEC)
_GOOD = RGBColor(0x2F, 0xA8, 0x4F)
_BAD = RGBColor(0xE2, 0x3B, 0x3B)
_PLACEHOLDER = RGBColor(0xE2, 0xE4, 0xE9)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# Cover accent "blobs"
_BLOBS = [
    (RGBColor(0x17, 0xBE, 0xB1), 62),   # teal
    (RGBColor(0xFF, 0xC2, 0x4B), 55),   # yellow
    (RGBColor(0xFF, 0x6B, 0x4A), 50),   # orange
    (RGBColor(0x7C, 0x5C, 0xFF), 58),   # violet
]

# Cost/efficiency metrics where a DECREASE is the good outcome — mirrors the
# frontend COST_METRICS set so delta colours carry the same GOOD/BAD semantic.
_COST_METRICS = frozenset({"cpa", "cpc", "cpm", "cpp", "frequency"})

_MAX_ADS = 6
_THUMB_TIMEOUT = 3.0
_INSIGHT_PLACEHOLDER = "Click to add your insight for this section…"

_CURRENCY_SYMBOLS = {
    "USD": "$", "CAD": "$", "AUD": "$", "SGD": "$",
    "EUR": "€", "GBP": "£", "JPY": "¥", "INR": "₹", "BRL": "R$",
}

# ─── Bundled brand assets (PNGs rasterized from frontend/public/*.svg) ────────────
# Kept as PNG because PowerPoint/python-pptx can't reliably embed SVG. A missing
# asset degrades to a text fallback — a logo is never allowed to fail the export.
_ASSET_DIR = os.path.join(os.path.dirname(__file__), "assets")
_PLATFORM_ICON = {"meta": "meta.png", "tiktok": "tiktok.png", "google_ads": "gads.png"}


def _asset(name: str) -> Optional[bytes]:
    # Read fresh each call: assets are small and swapping one on disk must take
    # effect without a process restart (no in-memory cache to go stale).
    try:
        with open(os.path.join(_ASSET_DIR, name), "rb") as fh:
            return fh.read()
    except Exception as exc:  # noqa: BLE001 - a missing asset degrades to text
        logger.info("export: brand asset %s unavailable: %s", name, exc)
        return None


# ─── Formatters (mirror frontend/src/lib/formatters.ts, server-side) ─────────────

def _fmt_money(value: Optional[float], currency: str = "USD", *, decimals: Optional[int] = None) -> str:
    """Format a currency value. `decimals=None` → 0 for >=1000 (keeps big values
    on one line, avoiding wrap/overlap), else 2. Unknown currencies are never
    converted (P-4) — the ISO code is prefixed instead of guessing a symbol."""
    if value is None:
        return "—"
    if decimals is None:
        decimals = 0 if abs(value) >= 1000 else 2
    body = f"{value:,.{decimals}f}"
    symbol = _CURRENCY_SYMBOLS.get(currency)
    return f"{symbol}{body}" if symbol else f"{currency} {body}"


def _fmt_number(value: Optional[float]) -> str:
    if value is None:
        return "—"
    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"{value / 1_000:.1f}K"
    return f"{int(round(value)):,}"


def _fmt_percent(value: Optional[float]) -> str:
    return "—" if value is None else f"{value:.2f}%"


def _fmt_roas(value: Optional[float]) -> str:
    return "—" if value is None else f"{value:.2f}x"


def _delta(pct: Optional[float], metric_key: str) -> tuple[str, RGBColor]:
    """(label, colour) for a period-over-period pct change. Direction encodes
    GOOD/BAD, not raw sign: for cost metrics a decrease is good (green)."""
    if pct is None:
        return "—", _MUTED
    if pct == 0:
        return "0.0%", _MUTED
    arrow = "▲" if pct > 0 else "▼"
    label = f"{arrow} {pct:+.1f}%"
    is_cost = metric_key in _COST_METRICS or metric_key.startswith("cost_per_")
    good = (pct < 0) if is_cost else (pct > 0)
    return label, (_GOOD if good else _BAD)


def _date_str(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value) if value is not None else ""


# ─── Primitives ──────────────────────────────────────────────────────────────────

def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, left, top, width, height, fill: RGBColor, *, shape=MSO_SHAPE.RECTANGLE,
          line: Optional[RGBColor] = None, line_w: float = 0.75, alpha: Optional[int] = None):
    sp = slide.shapes.add_shape(shape, Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if alpha is not None:
        _set_alpha(sp, alpha)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def _set_alpha(shape, opacity_pct: int) -> None:
    """Set solid-fill transparency (opacity 0..100) via the underlying XML —
    python-pptx has no public API for it."""
    spPr = shape._element.spPr
    solidFill = spPr.find(qn("a:solidFill"))
    if solidFill is None:
        return
    srgb = solidFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(opacity_pct * 1000))})
    srgb.append(a)


def _text(slide, left, top, width, height, text: str, *, size: int = 14, bold: bool = False,
          color: RGBColor = _INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          wrap: bool = True, italic: bool = False):
    box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def _rich(slide, left, top, width, height, parts, *, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, wrap: bool = False):
    """A single paragraph with multiple coloured runs. parts = [(text, size, bold, color)]."""
    box = slide.shapes.add_textbox(Emu(int(left)), Emu(int(top)), Emu(int(width)), Emu(int(height)))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{m}", 0)
    p = tf.paragraphs[0]
    p.alignment = align
    for txt, size, bold, color in parts:
        run = p.add_run()
        run.text = txt
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def _page_number(slide, n: int, *, dark: bool = False):
    _text(slide, _SLIDE_W - Emu(1500000), _SLIDE_H - Emu(430000), Emu(1200000), Emu(300000),
          f"PAGE  {n}", size=9, bold=True, color=(_MUTED if not dark else _WHITE), align=PP_ALIGN.RIGHT)


def _add_picture(slide, png: bytes, left, top, w, h):
    slide.shapes.add_picture(io.BytesIO(png), Emu(int(left)), Emu(int(top)), Emu(int(w)), Emu(int(h)))


def _wordmark(slide, *, dark: bool = False):
    """dashmet brand logo, top-right. Falls back to a text mark if the asset is
    missing so the export never depends on the image being present."""
    png = _asset("logo.png")
    size = 440000
    if png is not None:
        _add_picture(slide, png, _SLIDE_W - _MARGIN - size, Emu(300000), size, size)
    else:
        _text(slide, _SLIDE_W - Emu(2200000), Emu(300000), Emu(1900000), Emu(320000),
              "dashmet", size=13, bold=True, color=(_WHITE if dark else _INDIGO), align=PP_ALIGN.RIGHT)


def _platform_tag(slide, platform_id: str, left, top, *, dark: bool = False):
    """Platform logo + "<PLATFORM> ADS" label, left-aligned."""
    icon = _asset(_PLATFORM_ICON.get((platform_id or "").lower(), ""))
    x = left
    if icon is not None:
        d = 380000
        _add_picture(slide, icon, left, top, d, d)
        x = left + d + 150000
    _text(slide, x, top, Emu(5000000), Emu(380000),
          f"{(platform_id or '').upper()} ADS", size=12, bold=True,
          color=(_WHITE if dark else _INDIGO), anchor=MSO_ANCHOR.MIDDLE)


def _chevron_mark(slide, left, top):
    d = Emu(430000)
    ring = _rect(slide, left, top, d, d, _WHITE, shape=MSO_SHAPE.OVAL, line=_INDIGO, line_w=2.0)
    ring.shadow.inherit = False
    _text(slide, left, top - Emu(20000), d, d, "»", size=18, bold=True, color=_INDIGO,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _content_header(slide, title: str, ds: str, de: str, cds: str = "", cde: str = "") -> int:
    """Indigo section header + Date/Compare subline. Returns the y where content starts."""
    _chevron_mark(slide, _MARGIN, Emu(360000))
    _text(slide, _MARGIN + Emu(600000), Emu(360000), Emu(8000000), Emu(520000),
          title, size=26, bold=True, color=_INDIGO, anchor=MSO_ANCHOR.MIDDLE)
    _wordmark(slide)
    sub_y = Emu(960000)
    _text(slide, _MARGIN + Emu(600000), sub_y, Emu(5000000), Emu(320000),
          f"Date: {ds} – {de}", size=12, bold=True, color=_MUTED)
    if cds and cde:
        _text(slide, _SLIDE_W - Emu(6000000), sub_y, Emu(5440000), Emu(320000),
              f"Compare to: {cds} – {cde}", size=12, bold=True, color=_MUTED, align=PP_ALIGN.RIGHT)
    return 1400000


def _insight_box(slide, text: str = _INSIGHT_PLACEHOLDER, *, top=None, height=None):
    top = Emu(5560000) if top is None else Emu(int(top))
    height = Emu(880000) if height is None else Emu(int(height))
    box = _rect(slide, _MARGIN, top, _SLIDE_W - 2 * _MARGIN, height, _PURPLE_BOX,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(260000)
    tf.margin_right = Emu(260000)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = _WHITE


def _blobs_bg(slide):
    """Dark full-bleed background with a few translucent accent ovals (an
    approximation of the reference's fluid-gradient art)."""
    _rect(slide, 0, 0, _SLIDE_W, _SLIDE_H, _COVER_BG)
    # Spread so blobs bleed off distinct edges rather than muddying into one blot.
    placements = [
        (-1500000, -1400000, 3800000, 3800000, 0),   # teal   top-left
        (-1300000, 4200000, 3200000, 3200000, 1),     # yellow bottom-left
        (10600000, -1500000, 3000000, 3000000, 2),     # orange top-right
        (10200000, 4300000, 3400000, 3400000, 3),      # violet bottom-right
    ]
    for left, top, w, h, bi in placements:
        color, op = _BLOBS[bi]
        _rect(slide, left, top, w, h, color, shape=MSO_SHAPE.OVAL, alpha=op)


# ─── Thumbnail fetch (only network call; module-level so tests can patch) ────────

def _fetch_thumbnail(url: str) -> Optional[bytes]:
    try:
        resp = httpx.get(url, timeout=_THUMB_TIMEOUT, follow_redirects=True)
        if resp.status_code == 200:
            return resp.content
        logger.info("export: thumbnail HTTP %s for %s", resp.status_code, url)
        return None
    except Exception as exc:  # noqa: BLE001 - best-effort by design
        logger.info("export: thumbnail fetch failed for %s: %s", url, exc)
        return None


def _letterbox_45(raw: bytes) -> Optional[bytes]:
    """Contain-fit onto a 4:5 white canvas (object-contain parity). PNG bytes or None."""
    try:
        src = Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception as exc:  # noqa: BLE001
        logger.info("export: thumbnail decode failed: %s", exc)
        return None
    target_w, target_h = 400, 500
    canvas = Image.new("RGB", (target_w, target_h), (255, 255, 255))
    sw, sh = src.size
    if sw <= 0 or sh <= 0:
        return None
    scale = min(target_w / sw, target_h / sh)
    new_w, new_h = max(1, int(sw * scale)), max(1, int(sh * scale))
    resized = src.resize((new_w, new_h), Image.LANCZOS)
    canvas.paste(resized, ((target_w - new_w) // 2, (target_h - new_h) // 2))
    out = io.BytesIO()
    canvas.save(out, format="PNG")
    return out.getvalue()


# ─── Slides ──────────────────────────────────────────────────────────────────────

def _slide_cover(prs, account, overview: dict, exported_at: datetime) -> None:
    slide = _blank(prs)
    _blobs_bg(slide)
    period = overview.get("period") or {}
    ds, de = _date_str(period.get("date_start")), _date_str(period.get("date_stop"))
    preset = period.get("preset")

    _platform_tag(slide, account.platform_id, Emu(300000), Emu(320000), dark=True)
    _wordmark(slide, dark=True)

    _text(slide, Emu(300000), Emu(2350000), Emu(9000000), Emu(700000),
          "PERFORMANCE REPORT", size=44, bold=True, color=_WHITE)
    _text(slide, Emu(310000), Emu(3150000), Emu(9000000), Emu(560000),
          account.name, size=26, bold=True, color=RGBColor(0xFF, 0xC2, 0x4B))
    range_line = f"{ds} – {de}"
    if preset:
        range_line += f"  ({preset})"
    _text(slide, Emu(312000), Emu(3760000), Emu(9000000), Emu(360000),
          range_line, size=16, bold=True, color=_WHITE)
    _text(slide, Emu(312000), Emu(4160000), Emu(9000000), Emu(340000),
          f"{(account.platform_id or '').title()} · {account.currency} · {account.timezone}",
          size=13, color=RGBColor(0xB8, 0xBD, 0xC6))

    _text(slide, Emu(300000), _SLIDE_H - Emu(560000), Emu(9000000), Emu(320000),
          f"Data as of {de}   ·   Exported {exported_at.strftime('%Y-%m-%d %H:%M UTC')}",
          size=10, color=RGBColor(0x9A, 0x9F, 0xA8))


def _hero(slide, account, overview: dict, top: int) -> int:
    summary = overview.get("summary") or {}
    prev = overview.get("previous") or {}
    vs = overview.get("vs_previous") or {}
    left = _MARGIN
    width = _SLIDE_W - 2 * _MARGIN
    height = 1180000
    _rect(slide, left, top, width, height, _CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

    # Metric name + delta chip, centered near top
    dlabel, dcolor = _delta(vs.get("spend"), "spend")
    _rich(slide, left, top + 120000, width, 300000,
          [("Spend   ", 15, True, _INK), (dlabel, 12, True, dcolor)],
          align=PP_ALIGN.CENTER)

    third = width // 3
    # Current
    _text(slide, left, top + 520000, third, 260000, "Current", size=11, bold=True,
          color=_MUTED, align=PP_ALIGN.CENTER)
    _text(slide, left, top + 720000, third, 360000, _fmt_money(summary.get("spend"), account.currency),
          size=22, bold=True, color=_INK, align=PP_ALIGN.CENTER, wrap=False)
    # VS divider
    _text(slide, left + third, top + 620000, third, 400000, "VS", size=14, bold=True,
          color=_MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Previous
    _text(slide, left + 2 * third, top + 520000, third, 260000, "Previous", size=11, bold=True,
          color=_MUTED, align=PP_ALIGN.CENTER)
    _text(slide, left + 2 * third, top + 720000, third, 360000, _fmt_money(prev.get("spend"), account.currency),
          size=22, bold=True, color=_MUTED, align=PP_ALIGN.CENTER, wrap=False)
    return top + height + 220000


def _metric_card(slide, left, top, w, h, label, value, delta_label, delta_color, prev_str):
    _rect(slide, left, top, w, h, _CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pad = 200000
    _text(slide, left + pad, top + 150000, w - 2 * pad, 300000, label, size=12, bold=True, color=_MUTED)
    _text(slide, left + pad, top + 470000, w - 2 * pad, 360000, value, size=20, bold=True,
          color=_INK, wrap=False)
    tail = f"  vs {prev_str}" if prev_str else ""
    _rich(slide, left + pad, top + 860000, w - 2 * pad, 240000,
          [(delta_label, 11, True, delta_color), (tail, 11, False, _MUTED)])


def _slide_performance(prs, account, overview: dict, cds: str, cde: str) -> None:
    slide = _blank(prs)
    period = overview.get("period") or {}
    ds, de = _date_str(period.get("date_start")), _date_str(period.get("date_stop"))
    _content_header(slide, "MONTHLY PERFORMANCE", ds, de, cds, cde)

    summary = overview.get("summary") or {}
    prev = overview.get("previous") or {}
    vs = overview.get("vs_previous") or {}

    grid_top = _hero(slide, account, overview, 1420000)

    cur = account.currency
    cards = [
        ("Impressions", "impressions", _fmt_number),
        ("Clicks", "clicks", _fmt_number),
        ("CTR", "ctr", _fmt_percent),
        ("CPM", "cpm", lambda v: _fmt_money(v, cur)),
        ("Conversions", "conversions", _fmt_number),
        ("ROAS", "roas", _fmt_roas),
    ]
    cols, gap = 3, 200000
    usable = _SLIDE_W - 2 * _MARGIN
    cw = (usable - gap * (cols - 1)) // cols
    ch = 1160000
    rgap = 190000
    for i, (label, key, fmt) in enumerate(cards):
        r, c = divmod(i, cols)
        left = _MARGIN + c * (cw + gap)
        top = grid_top + r * (ch + rgap)
        dlabel, dcolor = _delta(vs.get(key), key)
        prev_str = fmt(prev.get(key)) if prev.get(key) is not None else ""
        _metric_card(slide, left, top, cw, ch, label, fmt(summary.get(key)), dlabel, dcolor, prev_str)

    _insight_box(slide, top=5560000, height=880000)
    _page_number(slide, 2)


def _series_points(series):
    return [p for p in (series or []) if p.get("spend") is not None]


def _slide_trend(prs, account, series: list[dict], previous_series, ds, de, cds, cde) -> None:
    slide = _blank(prs)
    _content_header(slide, f"DAILY CHART · Spend ({account.currency})", ds, de, cds, cde)

    cur_pts = _series_points(series)
    prev_pts = _series_points(previous_series)
    if not cur_pts:
        _text(slide, _MARGIN, Emu(1600000), _SLIDE_W - 2 * _MARGIN, Emu(600000),
              "No spend data for this period.", size=18, color=_MUTED)
        _page_number(slide, 3)
        return

    n = len(cur_pts)
    step = max(1, n // 8)  # thin x labels to ~8, fixing the crowded-axis bug
    cats = [(_date_str(p["date"]) if idx % step == 0 else "") for idx, p in enumerate(cur_pts)]

    chart_data = CategoryChartData()
    chart_data.categories = cats
    chart_data.add_series("Current", [float(p["spend"]) for p in cur_pts])
    if prev_pts:
        aligned = [float(prev_pts[i]["spend"]) if i < len(prev_pts) else None for i in range(n)]
        chart_data.add_series("Previous", aligned)

    gframe = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, _MARGIN, Emu(1450000), _SLIDE_W - 2 * _MARGIN, Emu(3850000), chart_data)
    chart = gframe.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.font.size = Pt(9)
    plot = chart.plots[0]
    # Current = indigo, Previous = muted amber.
    colors = [_INDIGO, RGBColor(0xF0, 0xB4, 0x29)]
    for si, ser in enumerate(plot.series):
        ser.format.line.color.rgb = colors[si % len(colors)]
        ser.format.line.width = Pt(2.25)
        ser.smooth = False
    cat_ax = chart.category_axis
    cat_ax.tick_labels.font.size = Pt(8)
    cat_ax.major_tick_mark = XL_TICK_MARK.NONE
    val_ax = chart.value_axis
    val_ax.tick_labels.font.size = Pt(8)
    val_ax.has_major_gridlines = True

    _insight_box(slide, top=5480000, height=980000)
    _page_number(slide, 3)


def _slide_campaigns(prs, account, overview: dict, ds, de, cds, cde) -> None:
    slide = _blank(prs)
    _content_header(slide, "TOP CAMPAIGNS", ds, de, cds, cde)

    campaigns = overview.get("top_campaigns") or []
    headers = ["Name", "Spend", "Impressions", "CTR", "Conversions", "ROAS"]
    if not campaigns:
        _text(slide, _MARGIN, Emu(1600000), _SLIDE_W - 2 * _MARGIN, Emu(600000),
              "No campaign data for this period.", size=18, color=_MUTED)
        _page_number(slide, 4)
        return

    nrows = len(campaigns) + 1
    width = _SLIDE_W - 2 * _MARGIN
    row_h = 560000
    height = row_h * nrows
    tbl = slide.shapes.add_table(nrows, len(headers), _MARGIN, Emu(1500000), width, Emu(height)).table
    tbl.first_row = False
    tbl.horz_banding = False

    ratios = [0.34, 0.15, 0.16, 0.10, 0.14, 0.11]
    for ci, ratio in enumerate(ratios):
        tbl.columns[ci].width = Emu(int(int(width) * ratio))
    for ri in range(nrows):
        tbl.rows[ri].height = Emu(row_h)

    def cell(r, c, text, *, bold=False, color=_INK, align=PP_ALIGN.LEFT, fill=None):
        cl = tbl.cell(r, c)
        cl.vertical_anchor = MSO_ANCHOR.MIDDLE
        cl.margin_top = Emu(40000)
        cl.margin_bottom = Emu(40000)
        cl.margin_left = Emu(140000)
        cl.margin_right = Emu(140000)
        if fill is not None:
            cl.fill.solid()
            cl.fill.fore_color.rgb = fill
        else:
            cl.fill.background()
        cl.text = ""
        p = cl.text_frame.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(12)
        run.font.bold = bold
        run.font.color.rgb = color

    for ci, h in enumerate(headers):
        cell(0, ci, h, bold=True, color=_WHITE, fill=_INDIGO,
             align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT)
    for ri, cmp in enumerate(campaigns, start=1):
        zebra = _CARD_ALT if ri % 2 else _CARD
        vals = [
            (cmp.get("name") or "—", PP_ALIGN.LEFT),
            (_fmt_money(cmp.get("spend"), account.currency), PP_ALIGN.RIGHT),
            (_fmt_number(cmp.get("impressions")), PP_ALIGN.RIGHT),
            (_fmt_percent(cmp.get("ctr")), PP_ALIGN.RIGHT),
            (_fmt_number(cmp.get("conversions")), PP_ALIGN.RIGHT),
            (_fmt_roas(cmp.get("roas")), PP_ALIGN.RIGHT),
        ]
        for ci, (val, align) in enumerate(vals):
            cell(ri, ci, val, align=align, fill=zebra, bold=(ci == 0))

    _insight_box(slide, top=5560000, height=880000)
    _page_number(slide, 4)


def _slide_ads(prs, account, ads: list[dict], ds, de, cds, cde) -> None:
    slide = _blank(prs)
    _content_header(slide, "TOP ADS BY SPEND", ds, de, cds, cde)

    top_ads = (ads or [])[:_MAX_ADS]
    if not top_ads:
        _text(slide, _MARGIN, Emu(1600000), _SLIDE_W - 2 * _MARGIN, Emu(600000),
              "No ad-level data for this period.", size=18, color=_MUTED)
        _page_number(slide, 5)
        return

    cols, rows_grid = 3, 2
    gap = 220000
    usable = _SLIDE_W - 2 * _MARGIN
    cell_w = (usable - gap * (cols - 1)) // cols
    top0 = 1480000
    cell_h = 2450000
    row_gap = 180000
    thumb_h = 1560000
    thumb_w = int(thumb_h * 4 / 5)

    for i, ad in enumerate(top_ads):
        r, c = divmod(i, cols)
        if r >= rows_grid:
            break
        left = _MARGIN + c * (cell_w + gap)
        top = top0 + r * (cell_h + row_gap)
        preview = ad.get("creative_preview") or {}
        title = preview.get("title") or ad.get("name") or "Untitled ad"
        url = preview.get("image_url") or preview.get("thumbnail_url")
        thumb_left = left + (cell_w - thumb_w) // 2

        png = None
        if url:
            try:
                raw = _fetch_thumbnail(url)
                if raw is not None:
                    png = _letterbox_45(raw)
            except Exception:  # noqa: BLE001 — degrade to placeholder, never fail the deck
                logger.warning("thumbnail render failed for ad %s", title, exc_info=True)
                png = None

        if png is not None:
            slide.shapes.add_picture(io.BytesIO(png), Emu(int(thumb_left)), Emu(int(top)),
                                     Emu(int(thumb_w)), Emu(int(thumb_h)))
        else:
            ph = _rect(slide, thumb_left, top, thumb_w, thumb_h, _PLACEHOLDER,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            ph.text_frame.word_wrap = True
            pp = ph.text_frame.paragraphs[0]
            pp.alignment = PP_ALIGN.CENTER
            run = pp.add_run()
            run.text = "No preview"
            run.font.size = Pt(11)
            run.font.color.rgb = _MUTED

        metrics = ad.get("metrics") or {}
        text_top = top + thumb_h + 70000
        _text(slide, left, text_top, cell_w, 340000, title, size=11, bold=True, color=_INK)
        stat = "   ·   ".join([
            _fmt_money(metrics.get("spend"), account.currency),
            f"CTR {_fmt_percent(metrics.get('ctr'))}",
            f"Conv {_fmt_number(metrics.get('conversions'))}",
            f"ROAS {_fmt_roas(metrics.get('roas'))}",
        ])
        _text(slide, left, text_top + 350000, cell_w, 300000, stat, size=9, color=_MUTED)

    _page_number(slide, 5)


def _slide_closing(prs, account) -> None:
    slide = _blank(prs)
    _blobs_bg(slide)
    logo = _asset("logo.png")
    if logo is not None:
        _add_picture(slide, logo, Emu(300000), Emu(1950000), Emu(560000), Emu(560000))
    _text(slide, Emu(300000), Emu(2680000), Emu(9000000), Emu(800000),
          "Thank You.", size=48, bold=True, color=_WHITE)
    _text(slide, Emu(312000), Emu(3630000), Emu(9000000), Emu(360000),
          f"Generated by dashmet · {account.name}", size=14, color=RGBColor(0xB8, 0xBD, 0xC6))
    _page_number(slide, 6, dark=True)


# ─── Public entrypoint ───────────────────────────────────────────────────────────

def generate_overview_pptx(
    *,
    account,
    overview: dict,
    series: list[dict],
    ads: list[dict],
    previous_series: Optional[list[dict]] = None,
    exported_at: Optional[datetime] = None,
) -> bytes:
    """Build the branded overview deck and return the .pptx as bytes.

    Args:
        account: SQLAlchemy Account model (name, platform_id, currency, timezone).
        overview: dict from insights.get_overview (period/summary/previous/vs_previous/top_campaigns).
        series: current-period timeseries points from get_timeseries(...)["series"].
        ads: ad-level rows from get_table(..., level="ad")[0] (already spend-desc).
        previous_series: prior-period timeseries points (get_timeseries(compare_previous=True)
            ["previous_series"]) for the trend overlay; None → single series.
        exported_at: optional stamp; defaults to now (UTC). Param keeps it testable.
    """
    if exported_at is None:
        exported_at = datetime.now(timezone.utc)

    period = overview.get("period") or {}
    ds, de = _date_str(period.get("date_start")), _date_str(period.get("date_stop"))
    # Compare window: derive from the previous_series if present (its span), else blank.
    cds = cde = ""
    prev_pts = _series_points(previous_series)
    if prev_pts:
        cds, cde = _date_str(prev_pts[0]["date"]), _date_str(prev_pts[-1]["date"])

    prs = Presentation()
    prs.slide_width = _SLIDE_W
    prs.slide_height = _SLIDE_H

    _slide_cover(prs, account, overview, exported_at)
    _slide_performance(prs, account, overview, cds, cde)
    _slide_trend(prs, account, series, previous_series, ds, de, cds, cde)
    _slide_campaigns(prs, account, overview, ds, de, cds, cde)
    _slide_ads(prs, account, ads, ds, de, cds, cde)
    _slide_closing(prs, account)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
