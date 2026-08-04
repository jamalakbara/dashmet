"""Tests for the PPTX overview-export feature (PRD §11).

Two layers:

  * Service-level (`generate_overview_pptx`) — no DB, no network. Builds the deck
    from synthetic overview/series/ads dicts shaped exactly like the shared
    insights read functions return. `export._fetch_thumbnail` is the ONLY network
    seam and is monkeypatched in every case so nothing leaves the process.
    Covers: a valid 6-slide deck, and thumbnail resilience (missing AND raising
    fetches must both still produce a 6-slide deck — P-4/P-8: a bad creative is a
    placeholder, never a failed export).

  * Endpoint-level (`overview_export_pptx`, DB-backed) — tenant isolation
    (cross-org → 403) and report-vs-API parity (P-6: the deck is fed the exact
    numbers `get_overview` returns, proven by capturing the `overview` kwarg the
    handler passes to `generate_overview_pptx` and comparing it to a direct
    `get_overview` call). These require a real Postgres (same as the rest of the
    suite); if it's unreachable they SKIP with a reason rather than error, so the
    service tests still gate the feature locally.

Run:
    cd backend && \
    $HOME/.pyenv/versions/3.12.13/bin/python -m pytest tests/test_export_overview.py -q
"""
import io
import types
import uuid
from datetime import date, datetime

import pytest
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from pptx import Presentation
from sqlalchemy import create_engine, text

from app.config import settings
from app.services import export

_PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


# ─── Synthetic account + read-path payloads (service layer, no DB) ────────────


def _fake_account(name="Acme Co", currency="USD", tz="UTC", platform="meta"):
    """A stand-in for the SQLAlchemy Account model — the exporter only reads
    .name/.platform_id/.currency/.timezone, so a namespace is sufficient and
    keeps this test DB-free."""
    return types.SimpleNamespace(
        name=name, platform_id=platform, currency=currency, timezone=tz,
    )


def _sample_overview():
    """Mirror of insights.get_overview(): period / summary / vs_previous /
    top_campaigns. Only the keys the exporter reads need to be present, but we
    populate the full summary the read path emits."""
    return {
        "period": {
            "date_start": date(2026, 6, 1),
            "date_stop": date(2026, 6, 30),
            "preset": "last_30d",
        },
        "summary": {
            "spend": 12345.67,
            "impressions": 2_500_000,
            "clicks": 48_000,
            "ctr": 1.92,
            "cpm": 4.94,
            "conversions": 1_234,
            "roas": 3.75,
        },
        "vs_previous": {
            "spend": 12.3,
            "impressions": -4.1,
            "clicks": 0.0,
            "ctr": None,
            "cpm": -2.5,
            "conversions": 8.8,
            "roas": 15.0,
        },
        "top_campaigns": [
            {
                "id": str(uuid.uuid4()), "name": "Alpha", "status": "active",
                "spend": 5000.0, "impressions": 1_000_000, "ctr": 2.1,
                "conversions": 600.0, "roas": 4.0, "outbound_clicks": 900.0,
            },
            {
                "id": str(uuid.uuid4()), "name": "Beta", "status": "paused",
                "spend": 2500.0, "impressions": 500_000, "ctr": 1.5,
                "conversions": 300.0, "roas": 3.0, "outbound_clicks": 400.0,
            },
        ],
    }


def _sample_series():
    return [
        {"date": date(2026, 6, 1), "spend": 400.0},
        {"date": date(2026, 6, 2), "spend": 415.5},
        {"date": date(2026, 6, 3), "spend": None},   # gap point must not crash
        {"date": date(2026, 6, 4), "spend": 500.0},
    ]


def _sample_ads(url="https://cdn.example.test/creative.jpg"):
    return [
        {
            "name": f"Ad {i}",
            "metrics": {"spend": 100.0 * i, "ctr": 1.5, "conversions": 10 * i, "roas": 2.0},
            "creative_preview": {
                "image_url": url, "thumbnail_url": url,
                "title": f"Creative {i}", "format": "IMAGE",
            },
        }
        for i in range(1, 4)
    ]


def _open_deck(raw: bytes) -> Presentation:
    assert isinstance(raw, bytes) and len(raw) > 0, "export produced empty bytes"
    return Presentation(io.BytesIO(raw))


def _deck_text(prs: Presentation) -> str:
    """All rendered text across every slide/shape, concatenated. Enough to assert
    a given string landed somewhere in a text box (the insight boxes render their
    text as a run in a rounded-rectangle shape)."""
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


# ─── 1. Service: valid deck (no DB, no network) ──────────────────────────────


def test_generate_overview_pptx_builds_six_slides(monkeypatch):
    """Happy path: a valid deck of exactly 6 slides, no network hit. Passing a
    previous_series exercises the current-vs-previous trend overlay branch."""
    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    prev_series = [
        {"date": date(2026, 5, 1), "spend": 380.0},
        {"date": date(2026, 5, 2), "spend": 402.0},
        {"date": date(2026, 5, 3), "spend": 450.0},
    ]
    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        previous_series=prev_series,
        ads=_sample_ads(),
        exported_at=datetime(2026, 8, 4, 9, 30),
    )

    prs = _open_deck(raw)
    assert len(prs.slides) == 6


# ─── 2. Service: thumbnail resilience (missing AND raising) ──────────────────


def test_generate_overview_pptx_survives_missing_thumbnail(monkeypatch):
    """A fetch that returns None (missing/404) degrades to a placeholder — the
    export still succeeds with 6 slides (P-4: no fake image, no failed export)."""
    calls = {"n": 0}

    def _none(url):
        calls["n"] += 1
        return None

    monkeypatch.setattr(export, "_fetch_thumbnail", _none)

    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        ads=_sample_ads(),
    )

    prs = _open_deck(raw)
    assert len(prs.slides) == 6
    assert calls["n"] >= 1, "ads with a preview url should attempt a fetch"


def test_generate_overview_pptx_survives_raising_thumbnail(monkeypatch):
    """A fetch that RAISES must not propagate: the export still yields 6 slides.

    `_fetch_thumbnail` swallows its own exceptions, but the ad-slide call site is
    ALSO guarded (export.py `_slide_ads`), so a raising replacement — a future
    refactor, an httpx/Pillow edge case — degrades to the neutral placeholder
    instead of taking down the whole deck (P-4/P-8: a bad creative is never a
    failed export)."""
    def _boom(url):
        raise RuntimeError("network exploded")

    monkeypatch.setattr(export, "_fetch_thumbnail", _boom)

    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        ads=_sample_ads(),
    )

    prs = _open_deck(raw)
    assert len(prs.slides) == 6


# ─── 2b. Service: AI insight-box injection (no DB, no network) ───────────────


def test_generate_overview_pptx_fills_insight_boxes(monkeypatch):
    """insights={...} threads each section's AI narrative into its insight box.

    The performance/trend/campaigns strings must appear in the rendered deck and
    the original placeholder must be GONE for those boxes (they were filled).
    This is the PPTX half of the AI-summary feature (P-6 deliverable-first)."""
    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    insights = {
        "performance": "PERF_NARRATIVE_MARKER performance moved up.",
        "trend": "TREND_NARRATIVE_MARKER momentum is improving.",
        "campaigns": "CAMPAIGNS_NARRATIVE_MARKER Alpha leads spend.",
    }
    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        ads=_sample_ads(),
        insights=insights,
    )

    text = _deck_text(_open_deck(raw))
    assert "PERF_NARRATIVE_MARKER performance moved up." in text
    assert "TREND_NARRATIVE_MARKER momentum is improving." in text
    assert "CAMPAIGNS_NARRATIVE_MARKER Alpha leads spend." in text
    # The three boxes that got AI text no longer show the placeholder. (The ads
    # slide has no insight box, so the placeholder should be entirely absent.)
    assert export._INSIGHT_PLACEHOLDER not in text


def test_generate_overview_pptx_partial_insights_keep_placeholder(monkeypatch):
    """A section missing from the insights dict falls back to the placeholder —
    filled boxes show AI text, unfilled boxes keep the original prompt (no
    layout change, graceful partial fill)."""
    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        ads=_sample_ads(),
        insights={"performance": "ONLY_PERF_MARKER"},
    )

    text = _deck_text(_open_deck(raw))
    assert "ONLY_PERF_MARKER" in text
    # trend + campaigns had no AI text → placeholder remains for those two boxes.
    assert export._INSIGHT_PLACEHOLDER in text


def test_generate_overview_pptx_none_insights_renders_placeholder(monkeypatch):
    """Regression guard: insights=None (default, token-free export) renders the
    original 'Click to add your insight…' placeholder exactly as before the AI
    feature — and none of the AI markers appear."""
    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    raw = export.generate_overview_pptx(
        account=_fake_account(),
        overview=_sample_overview(),
        series=_sample_series(),
        ads=_sample_ads(),
        insights=None,
    )

    text = _deck_text(_open_deck(raw))
    assert export._INSIGHT_PLACEHOLDER in text


# ─── DB availability guard (skip cleanly instead of erroring in fixtures) ─────


def _postgres_reachable() -> bool:
    try:
        engine = create_engine(settings.DATABASE_URL)
        with engine.connect() as conn:
            conn.execute(text("SELECT 1"))
        engine.dispose()
        return True
    except Exception:
        return False


requires_pg = pytest.mark.skipif(
    not _postgres_reachable(),
    reason="Postgres at settings.DATABASE_URL is unreachable; DB-backed export "
           "tests written but not executed (see conftest DB harness).",
)


# ─── 3. Endpoint: tenant isolation (DB) ──────────────────────────────────────


@requires_pg
def test_export_endpoint_cross_org_forbidden(db, monkeypatch):
    """A caller from org B exporting org A's account gets 403 — the handler maps
    ForbiddenError/NotFoundError from assert_account_belongs_to_org to
    HTTPException(403) inside _resolve_dates. Tenant isolation on every endpoint
    (PRD §11)."""
    from tests.conftest import (
        insert_metrics_daily, make_account, make_campaign,
        make_connection, make_org,
    )
    from app.api.v1.endpoints.insights import overview_export_pptx

    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    org_a = make_org(db, name="OrgA")
    conn_a = make_connection(db, org_a)
    acct_a = make_account(db, org_a, conn_a, name="A Account")
    camp_a = make_campaign(db, acct_a, name="A Camp")
    insert_metrics_daily(db, acct_a, camp_a, date(2026, 7, 20),
                         impressions=1000, clicks=50, spend=200)
    org_b = make_org(db, name="OrgB")
    db.flush()

    intruder = {"org_id": org_b, "sub": str(uuid.uuid4()), "role": "owner"}
    with pytest.raises(HTTPException) as ei:
        overview_export_pptx(
            account_id=acct_a,
            current_user=intruder,
            db=db,
            date_preset="last_30d",
            date_start=None,
            date_end=None,
            status=None,
            search=None,
        )
    assert ei.value.status_code == 403

    # The SAME call from the owning org returns a pptx StreamingResponse.
    owner = {"org_id": org_a, "sub": str(uuid.uuid4()), "role": "owner"}
    resp = overview_export_pptx(
        account_id=acct_a,
        current_user=owner,
        db=db,
        date_preset="last_30d",
        date_start=None,
        date_end=None,
        status=None,
        search=None,
    )
    assert isinstance(resp, StreamingResponse)
    assert resp.media_type == _PPTX_MEDIA_TYPE


# ─── 4. Report-vs-API parity (DB, P-6) ───────────────────────────────────────


@requires_pg
def test_export_summary_matches_get_overview(db, monkeypatch):
    """P-6: the deck is fed the EXACT summary that GET /overview returns — no
    divergent query path. Proven by capturing the `overview` kwarg the handler
    hands to generate_overview_pptx and comparing it to a direct get_overview
    call with the handler's own resolved date range."""
    from tests.conftest import (
        insert_metrics_daily, make_account, make_campaign,
        make_connection, make_org,
    )
    from app.api.v1.endpoints.insights import overview_export_pptx
    from app.services import export as export_svc_mod
    from app.services import insights as insights_svc

    org_id = make_org(db, name="ParityOrg")
    conn_id = make_connection(db, org_id)
    acct_id = make_account(db, org_id, conn_id, name="Parity Account", tz="UTC")
    camp_id = make_campaign(db, acct_id, name="Parity Camp")

    # Seed inside last_30d relative to today (2026-08-04 in the test env).
    ds, de = insights_svc.resolve_date_range("last_30d", "UTC")
    seed_day = de  # guaranteed within [ds, de]
    insert_metrics_daily(db, acct_id, camp_id, seed_day,
                         impressions=4000, clicks=120, spend=333.50)
    db.flush()

    # Direct read — the number of record the dashboard API would return.
    direct = insights_svc.get_overview(
        db, acct_id, org_id, ds, de, date_preset="last_30d",
    )

    # Capture the overview handed to the exporter (without building a real deck).
    captured = {}

    def _capture(*, account, overview, series, ads, exported_at=None):
        captured["overview"] = overview
        captured["series"] = series
        captured["ads"] = ads
        return b"stub-pptx-bytes"

    monkeypatch.setattr(export_svc_mod, "generate_overview_pptx", _capture)
    monkeypatch.setattr(export, "_fetch_thumbnail", lambda url: None)

    owner = {"org_id": org_id, "sub": str(uuid.uuid4()), "role": "owner"}
    resp = overview_export_pptx(
        account_id=acct_id,
        current_user=owner,
        db=db,
        date_preset="last_30d",
        date_start=None,
        date_end=None,
        status=None,
        search=None,
    )
    assert isinstance(resp, StreamingResponse)

    # Identical numbers by construction (P-6): same summary the API serves.
    assert captured["overview"]["summary"] == direct["summary"]
    # Sanity that the seeded spend actually flowed through, not two empty dicts.
    assert captured["overview"]["summary"]["spend"] == pytest.approx(333.50)
    assert captured["overview"]["summary"]["impressions"] == 4000
